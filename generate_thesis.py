# -*- coding: utf-8 -*-
"""生成毕设论文 .docx 文件"""

from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
import os, re

doc = Document()

# ── 页面设置 ──
for section in doc.sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)

# ── 样式设置 ──
style = doc.styles['Normal']
font = style.font
font.name = '宋体'
font.size = Pt(12)
style.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
pf = style.paragraph_format
pf.line_spacing = 1.5
pf.space_before = Pt(0)
pf.space_after = Pt(0)

def set_font(run, name='宋体', size=12, bold=False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor(*color)

def add_heading_custom(text, level=1):
    """添加标题，使用自定义字体"""
    if level == 0:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_font(run, '黑体', 22, bold=True)
        p.paragraph_format.space_before = Pt(30)
        p.paragraph_format.space_after = Pt(20)
        return p
    elif level == 1:
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_font(run, '黑体', 16, bold=True)
        p.paragraph_format.space_before = Pt(20)
        p.paragraph_format.space_after = Pt(12)
        return p
    elif level == 2:
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_font(run, '黑体', 14, bold=True)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(6)
        return p
    elif level == 3:
        p = doc.add_paragraph()
        run = p.add_run(text)
        set_font(run, '黑体', 12, bold=True)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        return p

def add_body(text):
    """添加正文段落，首行缩进2字符"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(text)
    set_font(run, '宋体', 12)
    return p

def add_mixed_paragraph(parts):
    """添加混合格式段落，parts = [(text, font_name, size, bold), ...]"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(24)
    p.paragraph_format.line_spacing = 1.5
    for text, name, size, bold in parts:
        run = p.add_run(text)
        set_font(run, name, size, bold)
    return p

def parse_md_content(text):
    """解析 Markdown 文本，返回段落列表"""
    paragraphs = []
    lines = text.split('\n')
    for line in lines:
        line = line.strip()
        if not line:
            continue
        if line.startswith('# ') or line.startswith('## ') or line.startswith('### '):
            level = line.count('#')
            title = line.lstrip('#').strip()
            paragraphs.append(('heading', title, level))
        elif line.startswith('| '):
            # 表格行，跳过简化处理
            continue
        elif line.startswith('- ') or line.startswith('1. '):
            paragraphs.append(('bullet', line, 0))
        else:
            paragraphs.append(('body', line, 0))
    return paragraphs

def add_from_md_file(filepath):
    """读取 .md 文件并添加到文档"""
    if not os.path.exists(filepath):
        add_body(f"[文件未找到: {filepath}]")
        return
    with open(filepath, 'r', encoding='utf-8') as f:
        content = f.read()
    paragraphs = parse_md_content(content)
    for ptype, text, level in paragraphs:
        if ptype == 'heading':
            add_heading_custom(text, level)
        elif ptype == 'body':
            add_body(text)
        elif ptype == 'bullet':
            add_body('• ' + text)

# ═══════════════════════════════════════════
# 封面
# ═══════════════════════════════════════════
for _ in range(4):
    doc.add_paragraph()

add_heading_custom('毕业论文', 0)
doc.add_paragraph()
add_heading_custom('DocAI——基于RAG与AI Agent的智能文档协作平台', 1)

for _ in range(3):
    doc.add_paragraph()

info_items = [
    ('学    院：', '_________________'),
    ('专    业：', '计算机科学与技术'),
    ('学生姓名：', '_________________'),
    ('学    号：', '_________________'),
    ('指导教师：', '_________________'),
]
for label, value in info_items:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(label)
    set_font(run, '宋体', 14)
    run = p.add_run(value)
    set_font(run, '宋体', 14)

doc.add_page_break()

# ═══════════════════════════════════════════
# 中文摘要
# ═══════════════════════════════════════════
add_heading_custom('摘  要', 1)

add_body('在大语言模型（Large Language Model, LLM）技术快速发展的背景下，企业文档管理正面临知识孤岛严重、AI 能力与业务场景脱节、协作效率低下以及重复性文档处理劳动密集等核心痛点。本文设计并实现了 DocAI——一个基于 RAG（Retrieval-Augmented Generation，检索增强生成）与 AI Agent 的智能文档协作平台。')

add_body('系统采用 Spring Boot 3.2 微服务架构，划分为网关服务、用户服务、文件服务、AI 服务和文档服务五个独立部署的模块，通过 Spring Cloud Gateway 统一入口、OpenFeign 声明式服务调用和 RabbitMQ 异步消息队列实现松耦合通信。在 AI 能力方面，系统基于 Spring AI 框架集成多种大语言模型，实现了完整的 RAG 检索链路——支持五种文档分段策略、三级检索体系和查询扩展机制，使大语言模型能够基于企业私有文档内容进行精准问答。同时，基于 ReAct 模式实现了 AI Agent 自主规划-执行-反思循环，注册了 17 种可调用工具。在协作方面，基于 STOMP WebSocket 协议和 OT 算法实现了多人实时编辑。此外，系统集成了 AIOps 运维监控模块。全部服务通过 Docker Compose 容器化编排，实现一键部署。')

add_body('实验结果表明，系统10项核心功能测试全部通过，AIOps 监控数据验证了各服务的稳定运行。本文工作验证了 RAG 与 AI Agent 技术在企业文档管理场景中的可行性和优越性。')

add_mixed_paragraph([
    ('关键词：', '黑体', 12, True),
    ('检索增强生成；AI Agent；ReAct 模式；微服务架构；实时协作；智能文档处理', '宋体', 12, False),
])

doc.add_page_break()

# ═══════════════════════════════════════════
# 英文摘要
# ═══════════════════════════════════════════
add_heading_custom('Abstract', 1)

add_body('Against the backdrop of rapid advancements in Large Language Models (LLMs), enterprise document management faces critical challenges including severe knowledge silos, disconnection between AI capabilities and business scenarios, low collaboration efficiency, and labor-intensive repetitive document processing tasks. This thesis presents DocAI, an intelligent document collaboration platform built upon Retrieval-Augmented Generation (RAG) and AI Agent technologies.')

add_body('The system adopts a Spring Boot 3.2 microservice architecture, comprising five independently deployable modules. Loose coupling is achieved through Spring Cloud Gateway for unified API entry, OpenFeign for declarative inter-service calls, and RabbitMQ for asynchronous message queuing. The platform implements a complete RAG pipeline featuring five document chunking strategies, a three-tier retrieval system, and a query expansion mechanism. Furthermore, the platform implements an AI Agent with a ReAct planning-execution-reflection loop, registering 17 callable tools. For real-time collaboration, the system supports multi-user co-editing through the STOMP WebSocket protocol and Operational Transformation algorithms. An AIOps monitoring module provides P95/P99 latency statistics with automatic fault detection. All services are containerized with Docker Compose for one-click deployment.')

add_body('Experimental results confirm that all 10 core functional tests pass and the AIOps monitoring data verifies stable service operation. This work validates the feasibility and superiority of combining RAG and AI Agent technologies in enterprise document management.')

add_mixed_paragraph([
    ('Keywords: ', 'Times New Roman', 12, True),
    ('Retrieval-Augmented Generation; AI Agent; ReAct Pattern; Microservice Architecture; Real-time Collaboration; Intelligent Document Processing', 'Times New Roman', 12, False),
])

doc.add_page_break()

# ═══════════════════════════════════════════
# 目录（占位）
# ═══════════════════════════════════════════
add_heading_custom('目  录', 1)
add_body('（请在 Word 中插入自动目录：引用 → 目录 → 自动目录）')
doc.add_page_break()

# ═══════════════════════════════════════════
# 章节内容
# ═══════════════════════════════════════════
base_dir = os.path.dirname(os.path.abspath(__file__))
chapter_files = [
    '第一章-绪论.md',
]

for ch in chapter_files:
    add_from_md_file(os.path.join(base_dir, ch))
    doc.add_page_break()

# 读取其他章节（在 F:\DocAI\ 下）
alt_dir = os.path.dirname(base_dir)  # F:\DocAI
alt_chapters = [
    '第二章-相关技术.md',
    '第三章-系统需求分析.md',
    '第四章-系统设计.md',
    '第五章-核心实现.md',
    '第六章-实验与验证.md',
    '第七章-总结与展望.md',
]
for ch in alt_chapters:
    fp = os.path.join(alt_dir, ch)
    if os.path.exists(fp):
        add_from_md_file(fp)
        doc.add_page_break()
    else:
        fp2 = os.path.join(base_dir, ch)
        if os.path.exists(fp2):
            add_from_md_file(fp2)
            doc.add_page_break()

# ═══════════════════════════════════════════
# 参考文献
# ═══════════════════════════════════════════
doc.add_page_break()
add_heading_custom('参考文献', 1)

refs = [
    '[1] LEWIS P, PEREZ E, PIKTUS A, et al. Retrieval-augmented generation for knowledge-intensive NLP tasks[C]//Advances in Neural Information Processing Systems 33 (NeurIPS 2020). Vancouver: Curran Associates, 2020: 9459-9474.',
    '[2] YAO S, ZHAO J, YU D, et al. ReAct: Synergizing reasoning and acting in language models[C]//The Eleventh International Conference on Learning Representations (ICLR 2023). Kigali, 2023.',
    '[3] KARPUKHIN V, OGUZ B, MIN S, et al. Dense passage retrieval for open-domain question answering[C]//Proceedings of the 2020 Conference on Empirical Methods in Natural Language Processing (EMNLP 2020). Online: ACL, 2020: 6769-6781.',
    '[4] SCHICK T, DWIVEDI-YU J, DESSI R, et al. Toolformer: Language models can teach themselves to use tools[C]//Advances in Neural Information Processing Systems 36 (NeurIPS 2023). New Orleans: Curran Associates, 2023.',
    '[5] XI Z, CHEN W, GUO X, et al. The rise and potential of large language model based agents: A survey[J]. arXiv preprint arXiv:2309.07864, 2023.',
    '[6] VASWANI A, SHAZEER N, PARMAR N, et al. Attention is all you need[C]//Advances in Neural Information Processing Systems 30 (NeurIPS 2017). Long Beach: Curran Associates, 2017: 5998-6008.',
    '[7] DEVLIN J, CHANG M W, LEE K, et al. BERT: Pre-training of deep bidirectional transformers for language understanding[C]//Proceedings of the 2019 Conference of the North American Chapter of the Association for Computational Linguistics (NAACL-HLT 2019). Minneapolis: ACL, 2019: 4171-4186.',
    '[8] BROWN T, MANN B, RYDER N, et al. Language models are few-shot learners[C]//Advances in Neural Information Processing Systems 33 (NeurIPS 2020). Vancouver: Curran Associates, 2020: 1877-1901.',
    '[9] ACHIAM J, ADLER S, AGARWAL S, et al. GPT-4 technical report[J]. arXiv preprint arXiv:2303.08774, 2023.',
    '[10] NEWMAN S. Building microservices: Designing fine-grained systems[M]. 2nd ed. Sebastopol: O\'Reilly Media, 2021.',
    '[11] ELLIS C A, GIBBS S J. Concurrency control in groupware systems[C]//Proceedings of the 1989 ACM SIGMOD International Conference on Management of Data (SIGMOD \'89). Portland: ACM, 1989: 399-407.',
    '[12] Spring Team. Spring Boot Reference Documentation (Version 3.2.0)[EB/OL]. (2023-11-23) [2024-06-01]. https://docs.spring.io/spring-boot/docs/3.2.0/reference/html/.',
    '[13] Spring Team. Spring AI Reference Documentation (Version 1.0.0-M4)[EB/OL]. (2024) [2024-06-01]. https://docs.spring.io/spring-ai/reference/.',
    '[14] Docker Inc. Docker Compose overview[EB/OL]. (2024) [2024-06-01]. https://docs.docker.com/compose/.',
    '[15] SHINN N, CASSANO F, GOPINATH A, et al. Reflexion: Language agents with verbal reinforcement learning[C]//Advances in Neural Information Processing Systems 36 (NeurIPS 2023). New Orleans: Curran Associates, 2023.',
    '[16] GAO Y, XIONG Y, GAO X, et al. Retrieval-augmented generation for large language models: A survey[J]. arXiv preprint arXiv:2312.10997, 2023.',
    '[17] BORGEAUD S, MENSCH A, HOFFMANN J, et al. Improving language models by retrieving from trillions of tokens[C]//International Conference on Machine Learning (ICML 2022). Baltimore: PMLR, 2022: 2206-2240.',
]

for ref in refs:
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(ref)
    set_font(run, '宋体', 10.5)

# ═══════════════════════════════════════════
# 致谢
# ═══════════════════════════════════════════
doc.add_page_break()
add_heading_custom('致  谢', 1)
add_body('（致谢内容待填写）')

# ═══════════════════════════════════════════
# 保存
# ═══════════════════════════════════════════
output_path = os.path.join(base_dir, '毕设论文-DocAI智能文档协作平台.docx')
doc.save(output_path)
print(f'论文已生成: {output_path}')
